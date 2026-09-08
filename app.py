import os
import re
import tempfile
from datetime import datetime
 
import streamlit as st
import pdfplumber
from docx import Document
 
from sop_analyzer import analyze_all
from explain_llm import enhance_explanation
 
# -------------------- PAGE CONFIG --------------------
st.set_page_config(
    page_title="Contract Risk Auditor",
    page_icon="⚖️",
    layout="wide",
)
 
# -------------------- MODEL PROVIDERS --------------------
# Every provider below speaks the OpenAI protocol, so only the endpoint and the
# model name change. Keeping them in one dict makes swapping providers a config
# change rather than a code change.
PROVIDERS = {
    "Zhipu GLM (free)": {
        "api_base": "https://open.bigmodel.cn/api/paas/v4",
        "model": "glm-4-flash",
        "key_hint": "open.bigmodel.cn",
    },
    "Google Gemini": {
        "api_base": "https://generativelanguage.googleapis.com/v1beta/openai/",
        "model": "gemini-2.0-flash",
        "key_hint": "aistudio.google.com",
    },
    "OpenAI": {
        "api_base": "https://api.openai.com/v1",
        "model": "gpt-4o-mini",
        "key_hint": "platform.openai.com",
    },
}
 
LEVEL_COLORS = {"RED": "#E5484D", "YELLOW": "#F5A524", "GREEN": "#30A46C"}
LEVEL_LABELS = {"RED": "High risk", "YELLOW": "Review needed", "GREEN": "Acceptable"}
 
# -------------------- STYLES --------------------
st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&family=JetBrains+Mono:wght@500&display=swap');
 
.stApp { background-color: #F4F5F7; }
html, body, [class*="css"], .stMarkdown, input, textarea, button, label {
    font-family: 'Inter', -apple-system, sans-serif !important;
}
 
/* Dark console header */
.console-bar {
    background: #14181F;
    border-radius: 10px;
    padding: 26px 32px;
    margin-bottom: 22px;
    display: flex;
    align-items: center;
    gap: 18px;
}
.console-bar h1 {
    color: #FFFFFF;
    font-size: 27px;
    font-weight: 700;
    letter-spacing: -0.4px;
    margin: 0;
}
.console-bar p {
    color: #8B93A1;
    font-size: 13.5px;
    margin: 4px 0 0 0;
    letter-spacing: 0.2px;
}
.console-tag {
    margin-left: auto;
    font-family: 'JetBrains Mono', monospace !important;
    font-size: 11px;
    color: #6E7787;
    border: 1px solid #2A303B;
    border-radius: 5px;
    padding: 5px 11px;
    white-space: nowrap;
}
 
/* Metric strip */
.metric-row { display: flex; gap: 12px; margin-bottom: 22px; }
.metric-box {
    flex: 1;
    background: #FFFFFF;
    border: 1px solid #E3E6EA;
    border-radius: 9px;
    padding: 16px 18px;
}
.metric-box .label {
    font-size: 11px;
    font-weight: 600;
    letter-spacing: 0.9px;
    text-transform: uppercase;
    color: #79808C;
}
.metric-box .value {
    font-size: 30px;
    font-weight: 700;
    line-height: 1.25;
    margin-top: 3px;
}
 
/* Clause cards */
.clause-card {
    background: #FFFFFF;
    border: 1px solid #E3E6EA;
    border-left-width: 4px;
    border-radius: 8px;
    padding: 17px 20px;
    margin-bottom: 13px;
}
.clause-head {
    display: flex;
    align-items: baseline;
    gap: 11px;
    margin-bottom: 12px;
}
.clause-title {
    font-family: 'JetBrains Mono', monospace !important;
    font-size: 13.5px;
    font-weight: 500;
    color: #14181F;
}
.clause-badge {
    margin-left: auto;
    font-size: 10.5px;
    font-weight: 700;
    letter-spacing: 0.7px;
    text-transform: uppercase;
    padding: 3px 9px;
    border-radius: 4px;
    color: #FFFFFF;
    white-space: nowrap;
}
.field-label {
    font-size: 10.5px;
    font-weight: 700;
    letter-spacing: 0.9px;
    text-transform: uppercase;
    color: #79808C;
    margin-top: 13px;
    margin-bottom: 4px;
}
.field-body { font-size: 14.2px; line-height: 1.62; color: #2B313B; }
.excerpt {
    font-size: 13.2px;
    line-height: 1.6;
    color: #4A515C;
    background: #F7F8FA;
    border-radius: 5px;
    padding: 10px 13px;
}
 
/* Drop zone */
.drop-zone {
    border: 1.5px dashed #C3C9D2;
    border-radius: 9px;
    padding: 30px;
    background: #FFFFFF;
    text-align: center;
    color: #79808C;
    font-size: 14px;
    margin-bottom: 14px;
}
[data-testid="stSidebar"] { background: #FFFFFF; }
</style>
""", unsafe_allow_html=True)
 
 
# -------------------- SESSION STATE --------------------
for key, default in [
    ("history", []),
    ("analysis", None),        # cached results for the current file
    ("analysis_key", None),    # file name + provider the cache belongs to
    ("api_key", ""),
]:
    if key not in st.session_state:
        st.session_state[key] = default
 
 
# -------------------- TEXT EXTRACTION --------------------
def extract_text(file) -> str:
    name = file.name.lower()
    if name.endswith(".txt"):
        return file.read().decode("utf-8", errors="ignore")
    if name.endswith(".docx"):
        return "\n".join(p.text for p in Document(file).paragraphs)
    if name.endswith(".pdf"):
        chunks = []
        with pdfplumber.open(file) as pdf:
            for page in pdf.pages:
                text = page.extract_text()
                if text:
                    chunks.append(text)
        return "\n".join(chunks)
    return ""
 
 
# -------------------- CLAUSE SPLITTING --------------------
def split_into_clauses(text):
    lines = text.splitlines()
    clauses = []
    current = {"title": None, "content": []}
    pattern = re.compile(r"^\s*(\d+(\.\d+)*)(\.|\))?\s")
    for line in lines:
        if pattern.match(line):
            if current["title"] and current["content"]:
                clauses.append({"title": current["title"], "text": "\n".join(current["content"]).strip()})
            current = {"title": line.strip(), "content": []}
        else:
            current["content"].append(line)
    if current["title"] and current["content"]:
        clauses.append({"title": current["title"], "text": "\n".join(current["content"]).strip()})
    return clauses
 
 
# -------------------- RISK SCORE --------------------
def risk_score(counters) -> int:
    """A single headline number. Deliberately simple and explainable:
    every high-risk clause costs 15 points, every review-needed clause costs 5."""
    penalty = counters.get("RED", 0) * 15 + counters.get("YELLOW", 0) * 5
    return max(0, 100 - penalty)
 
 
# -------------------- SIDEBAR --------------------
with st.sidebar:
    st.markdown("### Model provider")
    provider_name = st.selectbox("Provider", list(PROVIDERS.keys()), label_visibility="collapsed")
    st.caption(f"Get a key at {PROVIDERS[provider_name]['key_hint']}")
 
    st.markdown("### API key")
    api_key = st.text_input(
        "API key", type="password", label_visibility="collapsed",
        placeholder="Paste your API key...",
    )
    st.session_state.api_key = api_key
 
    st.markdown("---")
    st.markdown("### Recent files")
    if not st.session_state.history:
        st.caption("No files reviewed yet.")
    for fname in reversed(st.session_state.history[-6:]):
        st.caption(f"• {fname}")
 
 
# -------------------- HEADER --------------------
st.markdown("""
<div class="console-bar">
  <svg width="34" height="34" viewBox="0 0 24 24" fill="none" stroke="#FFFFFF"
       stroke-width="1.5" stroke-linecap="round" stroke-linejoin="round">
    <path d="M14 2H6a2 2 0 0 0-2 2v16a2 2 0 0 0 2 2h12a2 2 0 0 0 2-2V8z"></path>
    <polyline points="14 2 14 8 20 8"></polyline>
    <polyline points="9 15 11 17 15 13"></polyline>
  </svg>
  <div>
    <h1>Contract Risk Auditor</h1>
    <p>Rule-based clause screening for non-disclosure agreements</p>
  </div>
  <div class="console-tag">DETERMINISTIC ENGINE · 16 CHECKS</div>
</div>
""", unsafe_allow_html=True)
 
 
# -------------------- UPLOAD --------------------
st.markdown('<div class="drop-zone">Upload an NDA to run a full-document risk screen (PDF, DOCX or TXT)</div>',
            unsafe_allow_html=True)
uploaded = st.file_uploader("Upload contract", type=["pdf", "docx", "txt"], label_visibility="collapsed")
 
if not uploaded:
    st.stop()
 
if not api_key:
    st.warning("Enter an API key in the sidebar to run the review.")
    st.stop()
 
 
# -------------------- ANALYSIS (cached) --------------------
cache_key = f"{uploaded.name}::{provider_name}"
 
if st.session_state.analysis_key != cache_key:
    text = extract_text(uploaded)
 
    if not text.strip():
        st.error("No text could be extracted. This may be a scanned PDF, which would need OCR.")
        st.stop()
 
    clauses = split_into_clauses(text)
    if not clauses:
        st.error("No numbered clauses were found. This screen expects a clause-numbered contract.")
        st.stop()
 
    results, counters = analyze_all(clauses)
 
    # Only flagged clauses are sent to the model. Rewriting the explanation for
    # clauses that already passed adds cost and latency but no information.
    flagged = [r for r in results if r["level"] in ("RED", "YELLOW")]
    cfg = PROVIDERS[provider_name]
 
    if flagged:
        progress = st.progress(0.0, text=f"Reviewing {len(flagged)} flagged clauses...")
        for i, r in enumerate(flagged, start=1):
            r["risk"] = enhance_explanation(
                api_key, r["original_excerpt"], r["risk"], r["suggestion"],
                api_base=cfg["api_base"], model=cfg["model"],
            )
            progress.progress(i / len(flagged), text=f"Reviewing flagged clauses... {i}/{len(flagged)}")
        progress.empty()
 
    st.session_state.analysis = (results, counters)
    st.session_state.analysis_key = cache_key
    if uploaded.name not in st.session_state.history:
        st.session_state.history.append(uploaded.name)
 
results, counters = st.session_state.analysis
score = risk_score(counters)
score_color = "#E5484D" if score < 60 else "#F5A524" if score < 85 else "#30A46C"
 
 
# -------------------- DASHBOARD --------------------
st.markdown(f"""
<div class="metric-row">
  <div class="metric-box">
    <div class="label">Clauses screened</div>
    <div class="value" style="color:#14181F;">{len(results)}</div>
  </div>
  <div class="metric-box">
    <div class="label">High risk</div>
    <div class="value" style="color:{LEVEL_COLORS['RED']};">{counters.get('RED', 0)}</div>
  </div>
  <div class="metric-box">
    <div class="label">Review needed</div>
    <div class="value" style="color:{LEVEL_COLORS['YELLOW']};">{counters.get('YELLOW', 0)}</div>
  </div>
  <div class="metric-box">
    <div class="label">Acceptable</div>
    <div class="value" style="color:{LEVEL_COLORS['GREEN']};">{counters.get('GREEN', 0)}</div>
  </div>
  <div class="metric-box">
    <div class="label">Risk score</div>
    <div class="value" style="color:{score_color};">{score}<span style="font-size:15px;color:#79808C;">/100</span></div>
  </div>
</div>
""", unsafe_allow_html=True)
 
 
# -------------------- CLAUSE CARDS --------------------
def render_clause(r):
    color = LEVEL_COLORS[r["level"]]
    st.markdown(f"""
    <div class="clause-card" style="border-left-color:{color};">
      <div class="clause-head">
        <span class="clause-title">{r['title']}</span>
        <span class="clause-badge" style="background:{color};">{LEVEL_LABELS[r['level']]}</span>
      </div>
      <div class="field-label">Original excerpt</div>
      <div class="excerpt">{r['original_excerpt']}</div>
      <div class="field-label">Risk analysis</div>
      <div class="field-body">{r['risk']}</div>
      <div class="field-label">Suggested revision</div>
      <div class="field-body">{r['suggestion']}</div>
    </div>
    """, unsafe_allow_html=True)
 
 
tab_all, tab_red, tab_yellow, tab_green = st.tabs([
    f"All ({len(results)})",
    f"High risk ({counters.get('RED', 0)})",
    f"Review needed ({counters.get('YELLOW', 0)})",
    f"Acceptable ({counters.get('GREEN', 0)})",
])
 
with tab_all:
    for r in results:
        render_clause(r)
for tab, level in [(tab_red, "RED"), (tab_yellow, "YELLOW"), (tab_green, "GREEN")]:
    with tab:
        subset = [r for r in results if r["level"] == level]
        if not subset:
            st.caption("Nothing in this category.")
        for r in subset:
            render_clause(r)
 
 
# -------------------- REPORT --------------------
def save_report(results, counters, score, source_name):
    from docx.shared import RGBColor
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
 
    doc = Document()
    doc.add_heading("NDA Risk Assessment Report", level=1)
    doc.add_paragraph(f"Source document: {source_name}")
    doc.add_paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    doc.add_paragraph(
        f"Overall risk score: {score}/100  |  "
        f"High risk: {counters.get('RED', 0)}  |  "
        f"Review needed: {counters.get('YELLOW', 0)}  |  "
        f"Acceptable: {counters.get('GREEN', 0)}"
    )
    doc.add_paragraph("")
 
    for r in results:
        heading = doc.add_heading(level=2)
        run = heading.add_run(f"{r['title']} — {LEVEL_LABELS[r['level']]}")
        rgb = {"RED": (229, 72, 77), "YELLOW": (245, 165, 36), "GREEN": (48, 163, 108)}[r["level"]]
        run.font.color.rgb = RGBColor(*rgb)
 
        p = doc.add_paragraph()
        p.add_run("")
        p_format = p._p.get_or_add_pPr()
        border = OxmlElement("w:pBdr")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "18")
        left.set(qn("w:space"), "4")
        left.set(qn("w:color"), "%02X%02X%02X" % rgb)
        border.append(left)
        p_format.append(border)
 
        doc.add_paragraph("(1) Original (excerpt):")
        doc.add_paragraph(r["original_excerpt"])
        doc.add_paragraph("(2) Risk analysis:")
        doc.add_paragraph(r["risk"])
        doc.add_paragraph("(3) Suggested revision:")
        doc.add_paragraph(r["suggestion"])
        doc.add_paragraph("")
 
    path = os.path.join(tempfile.gettempdir(), "NDA_Risk_Report.docx")
    doc.save(path)
    return path
 
 
st.markdown("---")
report_path = save_report(results, counters, score, uploaded.name)
with open(report_path, "rb") as f:
    st.download_button(
        label="Download full report (.docx)",
        data=f.read(),
        file_name="NDA_Risk_Report.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
 









